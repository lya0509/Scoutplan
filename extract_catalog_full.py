import openpyxl
import json

wb = openpyxl.load_workbook("mi_plan_progresivo.xlsx", data_only=True)
sheet = wb["Masculino "]

catalog = []

current_super_area = ""
current_area = ""
current_sub_area = ""

for c in range(1, sheet.max_column + 1):
    r5 = str(sheet.cell(5, c).value or '').strip()
    r6 = str(sheet.cell(6, c).value or '').strip()
    r7 = str(sheet.cell(7, c).value or '').strip()
    r8 = str(sheet.cell(8, c).value or '').strip()

    if r5 and r5.isupper() and len(r5) > 3: current_super_area = r5
    if r6 and len(r6) > 3: current_area = r6
    if r7 and len(r7) > 3: current_sub_area = r7

    # Ignorar columnas de control/fechas/nombres de integrantes
    if not r8 or r8 in ["INTEGRANTES", "CARGO", "FECHA DE NACIMIENTO", "FECHA DE INGRESO (SCOUT)", "FECHA DE INGRESO (UNIDAD)", "FECHA DE PROMESA", "EDAD", "E", "F", "M", "A", "J", "S", "O", "N", "D", "FECHA DE ENTREGA", "VERIFICACIÓN"]:
        continue
    
    if len(r8) < 3:
        continue

    catalog.append({
        "col": c,
        "super_area": current_super_area,
        "area": current_area,
        "sub_area": current_sub_area,
        "activity": r8
    })

print(f"Total de actividades registradas en el catálogo oficial: {len(catalog)}")

with open("catalogo_completo_scout.json", "w", encoding="utf-8") as f:
    json.dump(catalog, f, ensure_ascii=False, indent=2)

print("Guardado en catalogo_completo_scout.json")

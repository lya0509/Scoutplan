import openpyxl
import json

wb = openpyxl.load_workbook("mi_plan_progresivo.xlsx", data_only=True)
sheet = wb["Masculino "]

col_data = []

for c in range(1, sheet.max_column + 1):
    r5 = str(sheet.cell(5, c).value or '').strip()
    r6 = str(sheet.cell(6, c).value or '').strip()
    r7 = str(sheet.cell(7, c).value or '').strip()
    r8 = str(sheet.cell(8, c).value or '').strip()
    
    if r5 or r6 or r7 or r8:
        col_data.append({
            "col": c,
            "r5": r5,
            "r6": r6,
            "r7": r7,
            "r8": r8
        })

print(f"Total de columnas procesadas con encabezados: {len(col_data)}")

with open("all_excel_columns.json", "w", encoding="utf-8") as f:
    json.dump(col_data, f, ensure_ascii=False, indent=2)

print("Guardado en all_excel_columns.json")

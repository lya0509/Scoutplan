import sys
import json

try:
    import openpyxl
except ImportError:
    print("El paquete 'openpyxl' no está instalado. Instalando openpyxl...")
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl"])
    import openpyxl

def process_excel(file_path):
    print(f"Cargando archivo: {file_path}...")
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
    except Exception as e:
        print(f"Error al abrir el archivo Excel: {e}")
        return

    print(f"¡Éxito! Se encontraron {len(wb.sheetnames)} hojas: {wb.sheetnames}")
    
    output_data = {}

    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        print(f"\n--- Procesando Hoja: {sheet_name} ---")
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue

        sheet_items = []
        for r_idx, row in enumerate(rows):
            row_items = [str(cell).strip() for cell in row if cell is not None and str(cell).strip() != '']
            if row_items:
                sheet_items.append({"row": r_idx + 1, "data": row_items})
        
        output_data[sheet_name] = sheet_items[:50] # muestra de resumen

    output_filename = "excel_summary.json"
    with open(output_filename, "w", encoding="utf-8") as f:
        json.dump(output_data, f, ensure_ascii=False, indent=2)

    print(f"\n✅ Procesamiento completado. Se generó un resumen en '{output_filename}'")

if __name__ == '__main__':
    if len(sys.argv) > 1:
        process_excel(sys.argv[1])
    else:
        print("Uso: py process_excel.py nombre_de_tu_archivo.xlsx")
